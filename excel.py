from datetime import datetime, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config
import db
from reports import period_bounds

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="305496")
MISSED_FILL = PatternFill("solid", fgColor="FCE4D6")
WRAP = Alignment(wrap_text=True, vertical="top")
KIND_NAMES = {"day": "Kunlik", "week": "Haftalik", "month": "Oylik"}


def _dt(ts: int | None) -> datetime | None:
    # Excel vaqt zonasini bilmaydi — mahalliy vaqtga o'tkazib, zonasiz yozamiz
    return datetime.fromtimestamp(ts, config.TZ).replace(tzinfo=None) if ts else None


def _person(name: str | None, username: str | None) -> str:
    return (name or "—") + (f" (@{username})" if username else "")


def _minutes(seconds: float | None) -> float | None:
    return round(seconds / 60, 1) if seconds is not None else None


def _link(chat_id: int, message_id: int) -> str | None:
    s = str(chat_id)
    return f"https://t.me/c/{s[4:]}/{message_id}" if s.startswith("-100") else None


def _sheet(wb: Workbook, title: str, headers: list[str], widths: list[int], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill, cell.alignment = HEADER_FONT, HEADER_FILL, WRAP
    ws.freeze_panes = "A2"
    return ws


def _finish(ws, date_cols: tuple[int, ...] = (), wrap_cols: tuple[int, ...] = ()) -> None:
    for row in ws.iter_rows(min_row=2):
        for c in date_cols:
            row[c - 1].number_format = "DD.MM.YYYY HH:MM"
        for c in wrap_cols:
            row[c - 1].alignment = WRAP
        link = row[-1]  # oxirgi ustun — Telegramdagi xabar havolasi
        if link.value:
            link.hyperlink, link.style = link.value, "Hyperlink"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def build_excel(chat_id: int, kind: str, previous: bool) -> tuple[bytes, str]:
    start, end, label = period_bounds(kind, previous)
    s_ts, e_ts = int(start.timestamp()), int(end.timestamp())
    stats = db.get_stats(chat_id, s_ts, e_ts)
    title = db.chat_title(chat_id)

    wb = Workbook()

    # 1. Umumiy — adminlar bo'yicha jamlanma
    ws = _sheet(wb, "Umumiy", ["№", "Admin", "Javob bergan savollari", "Jami javob xabarlari",
                               "O'rtacha javob vaqti (daq)"], [5, 35, 22, 22, 26], first=True)
    for i, a in enumerate(stats.admins, 1):
        ws.append([i, _person(a.name, a.username), a.answered, a.replies, _minutes(a.avg_response)])
    ws.append([])
    ws.append(["", "Guruh", title])
    ws.append(["", "Davr", label])
    ws.append(["", "Jami savollar", stats.total_questions])
    ws.append(["", "Javob berilgan", stats.answered])
    ws.append(["", "Javobsiz", stats.unanswered])
    ws.append(["", "O'rtacha javob vaqti (daq)", _minutes(stats.avg_response)])
    for row in ws.iter_rows(min_row=ws.max_row - 5, max_col=2):
        row[1].font = Font(bold=True)

    # 2. Savollar — har bir savol va unga kim javob bergani
    ws = _sheet(wb, "Savollar", [
        "№", "Savol vaqti", "Savol bergan", "Savol matni", "Holat", "Javob bergan admin",
        "Javob vaqti", "Javob muddati (daq)", "Javob matni", "Javoblar soni", "Havola",
    ], [5, 17, 25, 50, 15, 25, 17, 12, 50, 10, 35])
    for i, q in enumerate(db.get_questions(chat_id, s_ts, e_ts), 1):
        answered = q["answered_by"] is not None
        ws.append([
            i, _dt(q["created_at"]), _person(q["asker_name"], q["asker_username"]), q["text"],
            "Javob berilgan" if answered else "Javobsiz",
            _person(q["admin_name"], q["admin_username"]) if answered else None,
            _dt(q["answered_at"]),
            _minutes(q["answered_at"] - q["created_at"]) if answered else None,
            q["answer_text"], q["reply_count"], _link(chat_id, q["message_id"]),
        ])
        if not answered:
            ws.cell(row=ws.max_row, column=5).fill = MISSED_FILL
    _finish(ws, date_cols=(2, 7), wrap_cols=(4, 9))

    # 3. Admin javoblari — har bir admin yozgan barcha javoblar
    ws = _sheet(wb, "Admin javoblari", [
        "№", "Admin", "Javob vaqti", "Javob matni", "Kimga", "Savol vaqti", "Savol matni", "Havola",
    ], [5, 25, 17, 50, 25, 17, 50, 35])
    for i, r in enumerate(db.get_replies(chat_id, s_ts, e_ts), 1):
        ws.append([
            i, _person(r["admin_name"], r["admin_username"]), _dt(r["created_at"]), r["text"],
            _person(r["asker_name"], r["asker_username"]) if r["q_created_at"] else "(savol sifatida qayd etilmagan)",
            _dt(r["q_created_at"]), r["q_text"], _link(chat_id, r["message_id"]),
        ])
    _finish(ws, date_cols=(3, 6), wrap_cols=(4, 7))

    buf = BytesIO()
    wb.save(buf)
    filename = f"{KIND_NAMES[kind]}_{start:%d.%m.%Y}"
    if kind != "day":
        filename += f"-{end - timedelta(seconds=1):%d.%m.%Y}"
    return buf.getvalue(), filename + ".xlsx"
